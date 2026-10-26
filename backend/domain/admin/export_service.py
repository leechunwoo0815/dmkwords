# backend/domain/admin/export_service.py — 管理端 Excel 导出（从 service.py 拆出，P3 性能族）
"""为什么拆出来（2026-10-22 第一百一十五批）：`admin/service.py` 已贴 800 行 god-file 红线
（交接卡 §六 早就登记「下次动它先把导出三件套挪出去」），本批给仪表盘加 TTL 缓存时正好越线。
拆法与 `task_specs.py` 同款：**原模块保持 re-export**，router 等既有 import 路径不变。

口径不变：两个导出都是"上限在文件里如实写明，不静默截断"（P2-44），审计导出仍走
`safe_row` 净化（P2-18：低权限可写文本不得注入超管的 Excel 公式）。
"""

from __future__ import annotations

from sqlalchemy.orm import Session

from backend.domain.admin.models import AuditLog


class AuditExportService:
    """审计日志 Excel 导出（C18：FEAT-005「可查询导出」）。"""

    def __init__(self, db: Session):
        self.db = db

    def export_excel(self) -> bytes:
        """审计日志导出（P2-44：上限在文件里如实写明，不再静默截断）。"""
        from io import BytesIO

        from openpyxl import Workbook
        from openpyxl.styles import Font

        from backend.domain.admin.audit_labels import AUDIT_ACTION_LABEL, AUDIT_TARGET_LABEL
        from backend.domain.admin.service import AUDIT_EXPORT_LIMIT

        q = self.db.query(AuditLog).filter(AuditLog.is_deleted == 0)
        total = q.count()
        rows = q.order_by(AuditLog.id.desc()).limit(AUDIT_EXPORT_LIMIT).all()
        wb = Workbook()
        ws = wb.active
        ws.title = "操作审计"
        headers = ["ID", "操作人", "动作", "对象类型", "对象ID", "详情", "原因", "时间"]
        if total > len(rows):
            ws.append(
                [
                    f"注意：符合条件的记录共 {total} 条，本表仅导出最新 {len(rows)} 条"
                    f"（上限 {AUDIT_EXPORT_LIMIT}）——请缩小筛选范围后分批导出。"
                ]
            )
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        from backend.common.sheet_safety import safe_row

        for r in rows:
            # P2-18（审查点名处）：审计 `reason` 是**专员可写**的自由文本，原样落进超管导出的
            # Excel 就是"低权限对高权限工作站的公式注入"——统一净化后再写。
            ws.append(
                safe_row(
                    [
                        r.id,
                        r.actor_name,
                        AUDIT_ACTION_LABEL.get(r.action, r.action),
                        AUDIT_TARGET_LABEL.get(r.target_type, r.target_type),
                        r.target_id,
                        r.detail or "",
                        r.reason,
                        r.created_at.strftime("%Y-%m-%d %H:%M:%S") if r.created_at else "",
                    ]
                )
            )
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()


class DashboardExportService:
    """数据看板 Excel 导出（docs/04 WM11 步骤 7：导出任意报表 Excel）。"""

    def __init__(self, db: Session):
        self.db = db

    def export_excel(self) -> bytes:
        from io import BytesIO

        from openpyxl import Workbook
        from openpyxl.styles import Font

        from backend.domain.admin.service import DashboardService

        data = DashboardService(self.db).overview()
        wb = Workbook()
        ws = wb.active
        ws.title = "数据看板"
        ws.append(["指标", "数值"])
        for c in ws[1]:
            c.font = Font(bold=True)
        label_map = {
            "admin_count": "后台账号",
            "today_logins": "今日登录",
            "config_count": "业务配置项",
            "copy_total": "总藏书量",
            "copy_available": "在馆",
            "copy_borrowed": "借出",
            "copy_maintenance": "维护",
            "copy_lost": "遗失",
            "today_borrowed": "今日借出",
            "today_returned": "今日归还",
            "overdue_active": "当前逾期",
            "member_total": "会员总数",
            "member_new_week": "本周新增会员",
            "pending_evaluation_count": "待评估人数",
            "activity_enroll_recent": "近7天活动报名",
            "renew_rate": "续借率(%)",
            "withdrawal_rate": "退会率(%)",
            "quiz_pass_rate": "测验通过率(%)",
            "milestone_count": "里程碑达成人数",
        }
        for key, label in label_map.items():
            ws.append([label, data.get(key, 0)])
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
